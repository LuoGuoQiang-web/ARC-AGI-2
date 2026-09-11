# -*- coding: utf-8 -*-
import io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import docx

d = docx.Document(r'C:\Users\LRDC07\Desktop\Kaggle\铝合金材料性能预测建模专属化方案.docx')

for p in d.paragraphs:
    print(f'[{p.style.name}] {p.text}')

print('---TABLES---')
for ti, t in enumerate(d.tables):
    print(f'== table {ti} ==')
    for r in t.rows:
        print(' | '.join(c.text.replace('\n', ' ') for c in r.cells))
